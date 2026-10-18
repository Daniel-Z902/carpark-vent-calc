"""Generate extra golden cases by recalculating Daniel's own sheets with LibreOffice.

The workbook ships with example data only on the 2012, 2002 and 1991 sheets
(and every queue result is 0).  To cover the 2024 sheet and non-zero queue
lengths, this script writes synthetic inputs into the *unchanged* formula
rows of carpark_calcs.xlsx, recalculates headless with LibreOffice and stores
inputs + recalculated results in tests/golden/libreoffice_recalc.json.

    python tools/make_lo_scenarios.py [carpark_calcs.xlsx]

Needs `soffice` on PATH.  Results are LibreOffice's evaluation of the Excel
formulas, not Excel cached values - tests label them accordingly.
"""

from __future__ import annotations

import json
import random
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

from carpark_vent.editions import EDITIONS
from carpark_vent.layout import LAYOUTS

SEED = 1668


def _scenarios(key: str, rng: random.Random) -> list[dict]:
    ed = EDITIONS[key]
    lay = LAYOUTS[key]
    nrows = lay.last_data_row - lay.first_data_row + 1
    usage = list(ed.usage_factors.values())
    vehicle = list(ed.vehicle_factors.values()) or [1.0]
    out = []
    for sid, (lanes, heavy) in enumerate([(1, False), (3, False), (1, True), (2, True)]):
        levels = []
        for i in range(nrows):
            lv = {"name": f"L{i + 1}", "n1": rng.randint(400, 900) if heavy else rng.randint(1, 400)}
            for f in ed.level_fields:
                if f == "n1":
                    continue
                if f == "n2":
                    lv[f] = rng.choice([0, rng.randint(1, 800)])
                elif f.startswith("d"):
                    lv[f] = round(rng.uniform(0, 120), 2)
                elif f in ("P", "P1", "P2"):
                    lv[f] = rng.choice(usage[2:] if heavy else usage)
                elif f == "T":
                    lv[f] = rng.choice(vehicle)
                elif f in ("E", "F"):
                    lv[f] = rng.choice([0.5, 1, 1.5, 2])
                elif f == "A":
                    lv[f] = round(rng.uniform(300, 9000), 2)
            levels.append(lv)
        out.append({"id": f"lo-{key}-{sid}", "edition": key, "exit_lanes": lanes, "levels": levels})
    return out


def main() -> None:
    src = Path(sys.argv[1] if len(sys.argv) > 1 else "carpark_calcs.xlsx").resolve()
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("soffice not found")
    rng = random.Random(SEED)
    formulas = openpyxl.load_workbook(src)
    cases = []
    with tempfile.TemporaryDirectory() as tmp:
        tmpd = Path(tmp)
        files = []
        for key in EDITIONS:
            for sc in _scenarios(key, rng):
                wb = openpyxl.load_workbook(src)
                ws = wb[EDITIONS[key].sheet_name]
                lay = LAYOUTS[key]
                for r in range(lay.first_data_row, lay.last_data_row + 1):
                    for col in lay.inputs.values():
                        ws[f"{col}{r}"] = None
                for i, lv in enumerate(sc["levels"]):
                    r = lay.first_data_row + i
                    for f, val in lv.items():
                        ws[f"{lay.inputs[f]}{r}"] = val
                ws[lay.exit_lanes_cell] = sc["exit_lanes"]
                fn = tmpd / f"{sc['id']}.xlsx"
                wb.save(fn)
                files.append((sc, fn))
        outdir = tmpd / "out"
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(outdir),
                        *[str(f) for _, f in files]], check=True, capture_output=True, timeout=600)
        for sc, fn in files:
            key = sc["edition"]
            lay = LAYOUTS[key]
            fws = formulas[EDITIONS[key].sheet_name]
            v = openpyxl.load_workbook(outdir / fn.name, data_only=True)[EDITIONS[key].sheet_name]
            exp = []
            for i in range(len(sc["levels"])):
                r = lay.first_data_row + i
                row = {k: v[f"{c}{r}"].value for k, c in lay.outputs.items()}
                row["queue_cells_are_formulas"] = all(
                    str(fws[f"{lay.outputs[k]}{r}"].value).startswith("=") for k in ("lb_lhu", "llu"))
                area_col = lay.outputs["Qd" if key == "1991" else "Qc"]
                m = re.search(r"([\d.]+)\*[A-Z]+\d+\)$", str(fws[f"{area_col}{r}"].value))
                row["sheet_area_rate"] = float(m.group(1)) if m else None
                row["row"] = r
                exp.append(row)
            totals = {k: v[f"{c}{lay.total_row}"].value for k, c in lay.totals.items()}
            cases.append({
                "id": sc["id"],
                "source": f"LibreOffice recalculation of sheet '{EDITIONS[key].sheet_name}' with synthetic inputs",
                "input": {"edition": key, "exit_lanes": sc["exit_lanes"], "levels": sc["levels"]},
                "expected": {"levels": exp, "totals": totals},
            })
    out = Path("tests/golden/libreoffice_recalc.json")
    out.write_text(json.dumps({"workbook": src.name, "seed": SEED, "cases": cases}, indent=1) + "\n")
    print(f"wrote {out}: {len(cases)} cases")


if __name__ == "__main__":
    main()
