"""Report writers and CLI smoke tests, plus a live-formula round trip through LibreOffice."""

import json
import shutil
import subprocess

import openpyxl
import pytest

from carpark_vent import calculate_project
from carpark_vent.cli import main
from carpark_vent.layout import LAYOUTS
from carpark_vent.loader import load_project, project_from_dict
from carpark_vent.report import render_html, write_excel

from conftest import ROOT, close, load_cases

EXAMPLES = sorted((ROOT / "examples").glob("*.*"))


@pytest.mark.parametrize("path", EXAMPLES, ids=[p.name for p in EXAMPLES])
def test_cli_run_all_formats(path, tmp_path, capsys):
    assert main(["run", str(path), "-o", str(tmp_path)]) == 0
    for ext in ("xlsx", "html", "json", "pdf"):
        assert (tmp_path / f"{path.stem}.{ext}").stat().st_size > 0
    data = json.loads((tmp_path / f"{path.stem}.json").read_text())
    res = calculate_project(load_project(path))
    assert close(data["totals"]["design_exhaust_m3s"], res.total_design_exhaust_m3s)
    assert "TOTAL" in capsys.readouterr().out


def test_example_2012_reproduces_workbook_totals():
    res = calculate_project(load_project(ROOT / "examples" / "example_2012.yaml"))
    case = next(c for c in load_cases("excel_cached.json") if c["id"] == "excel-2012")
    assert close(res.total_design_exhaust_m3s, case["expected"]["totals"]["design"])
    assert close(res.total_makeup_m3s, case["expected"]["totals"]["makeup"])


def test_template_command_round_trips(tmp_path, capsys):
    assert main(["template"]) == 0
    default_text = capsys.readouterr().out
    assert 'edition: "2012"' in default_text
    default_file = tmp_path / "t-default.yaml"
    default_file.write_text(default_text)
    assert load_project(default_file).edition_key == "2012"
    for ed in ("2012", "2024"):
        assert main(["template", ed]) == 0
        f = tmp_path / f"t{ed}.yaml"
        f.write_text(capsys.readouterr().out)
        assert f'edition: "{ed}"' in f.read_text()
        calculate_project(load_project(f))


def test_cli_reports_input_errors(tmp_path, capsys):
    bad = tmp_path / "bad.json"
    bad.write_text(json.dumps({"edition": "2012", "levels": [{"name": "x", "n1": 1}]}))
    assert main(["run", str(bad), "-o", str(tmp_path)]) == 2
    assert "'P' is required" in capsys.readouterr().err


def test_excel_values_mirror_layout(tmp_path):
    res = calculate_project(load_project(ROOT / "examples" / "example_2012.yaml"))
    p = write_excel(res, tmp_path / "r.xlsx")
    ws = openpyxl.load_workbook(p).worksheets[0]
    lay = LAYOUTS["2012"]
    assert ws["A34"].value == "BASEMENT 2" and ws["B7"].value == 2000
    assert ws[f"{lay.outputs['design']}35"].value == pytest.approx(res.levels[1].design_exhaust_m3s)
    assert ws[f"{lay.totals['design']}36"].value == pytest.approx(res.total_design_exhaust_m3s)
    assert ws[lay.makeup_ratio_cell].value == 0.8


def test_html_contains_sections():
    res = calculate_project(load_project(ROOT / "examples" / "example_2024.yaml"))
    h = render_html(res)
    for s in ("Project details", "Method", "Inputs", "Step-by-step", "Summary", "AS 1668.2-2024"):
        assert s in h


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
@pytest.mark.parametrize("case_id", ["lo-2024-2", "lo-2012-2"])
def test_live_formula_workbook_recalculates_to_python_values(case_id, tmp_path):
    case = next(c for c in load_cases("libreoffice_recalc.json") if c["id"] == case_id)
    res = calculate_project(project_from_dict(case["input"]))
    src = write_excel(res, tmp_path / "f.xlsx", formulas=True)
    out = tmp_path / "out"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=180)
    ws = openpyxl.load_workbook(out / "f.xlsx", data_only=True).worksheets[0]
    lay = LAYOUTS[case["input"]["edition"]]
    for i, lr in enumerate(res.levels):
        r = lay.first_data_row + i
        assert close(ws[f"{lay.outputs['design']}{r}"].value, lr.design_exhaust_m3s)
        assert close(ws[f"{lay.outputs['makeup']}{r}"].value, lr.makeup_m3s)
        assert close(ws[f"{lay.outputs['queue']}{r}"].value, lr.queue_exhaust_m3s)
    tr = lay.first_data_row + len(res.levels)
    assert close(ws[f"{lay.totals['design']}{tr}"].value, res.total_design_exhaust_m3s)
